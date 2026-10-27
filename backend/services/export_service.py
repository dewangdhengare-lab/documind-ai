"""Export service for generating PDF and DOCX reports."""

from __future__ import annotations

import io
import json
import logging
import re
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors

logger = logging.getLogger(__name__)


def export_report_to_docx(title: str, content: str) -> io.BytesIO:
    """Convert markdown content to a formatted DOCX document."""
    doc = DocxDocument()

    # Document Title
    title_p = doc.add_heading(title, level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    lines = content.split("\n")
    in_table = False
    table_lines: list[str] = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        rows = []
        for line in table_lines:
            if re.match(r"^\|?\s*[-:]+[-| :]*$", line.strip()):
                continue  # Divider line
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if cells:
                rows.append(cells)

        if rows:
            col_count = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=col_count)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx, cell_text in enumerate(row):
                    if c_idx < col_count:
                        table.cell(r_idx, c_idx).text = cell_text

        table_lines = []
        in_table = False
        doc.add_paragraph()

    for line in lines:
        stripped = line.strip()

        # Check for table line
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_markdown_runs_to_paragraph(p, stripped[2:])
        elif re.match(r"^\d+\.\s+", stripped):
            clean_text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            _add_markdown_runs_to_paragraph(p, clean_text)
        else:
            p = doc.add_paragraph()
            _add_markdown_runs_to_paragraph(p, stripped)

    if in_table:
        flush_table()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_markdown_runs_to_paragraph(paragraph, text: str):
    """Simple parser to preserve bold **text** in DOCX paragraphs."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def export_report_to_pdf(title: str, content: str) -> io.BytesIO:
    """Convert markdown content to a formatted PDF document using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#312E81"),
        spaceAfter=18,
    )
    h1_style = ParagraphStyle(
        "DocH1",
        parent=styles["Heading1"],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#4338CA"),
        spaceBefore=12,
        spaceAfter=8,
        keepWithNext=True,
    )
    h2_style = ParagraphStyle(
        "DocH2",
        parent=styles["Heading2"],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#4F46E5"),
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#1F2937"),
        leftIndent=15,
        spaceAfter=4,
    )

    story = [Paragraph(title, title_style), Spacer(1, 12)]

    lines = content.split("\n")
    in_table = False
    table_lines: list[str] = []

    def flush_pdf_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return

        rows = []
        for line in table_lines:
            if re.match(r"^\|?\s*[-:]+[-| :]*$", line.strip()):
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if cells:
                rows.append([Paragraph(_convert_markdown_inline(c), body_style) for c in cells])

        if rows:
            col_count = max(len(r) for r in rows)
            # Normalize row cell count
            for r in rows:
                while len(r) < col_count:
                    r.append(Paragraph("", body_style))

            t = Table(rows)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2FF")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

        table_lines = []
        in_table = False

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_pdf_table()

        if not stripped:
            continue

        if stripped.startswith("### "):
            story.append(Paragraph(stripped[4:], h2_style))
        elif stripped.startswith("## "):
            story.append(Paragraph(stripped[3:], h1_style))
        elif stripped.startswith("# "):
            story.append(Paragraph(stripped[2:], title_style))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            formatted = "&bull; " + _convert_markdown_inline(stripped[2:])
            story.append(Paragraph(formatted, bullet_style))
        elif re.match(r"^\d+\.\s+", stripped):
            clean_text = _convert_markdown_inline(stripped)
            story.append(Paragraph(clean_text, bullet_style))
        else:
            story.append(Paragraph(_convert_markdown_inline(stripped), body_style))

    if in_table:
        flush_pdf_table()

    doc.build(story)
    buffer.seek(0)
    return buffer


def _convert_markdown_inline(text: str) -> str:
    """Convert inline markdown (bold, italic, code) to ReportLab XML tags."""
    # Escape XML entities first
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    # Bold **text**
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    # Italic *text*
    text = re.sub(r"\*(.*?)\*", r"<i>\1</i>", text)
    # Inline code `text`
    text = re.sub(r"`(.*?)`", r"<font face='Courier'>\1</font>", text)
    return text
